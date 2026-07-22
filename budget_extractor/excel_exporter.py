"""Professionally formatted Excel workbook export."""

from __future__ import annotations

from datetime import datetime, timezone
from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Alignment, Font, PatternFill, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

from budget_extractor.logging_utils import RunLogEntry
from budget_extractor.schemas import FinalExtractionOutput
from budget_extractor.utils import document_stem, ensure_dir, flatten_list

HEADER_FILL = PatternFill("solid", fgColor="1F4E79")
HEADER_FONT = Font(bold=True, color="FFFFFF")
ALT_ROW_FILL = PatternFill("solid", fgColor="F2F2F2")
GREEN_FILL = PatternFill("solid", fgColor="C6EFCE")
YELLOW_FILL = PatternFill("solid", fgColor="FFEB9C")
RED_FILL = PatternFill("solid", fgColor="FFC7CE")
THIN = Border(
    left=Side(style="thin", color="D9D9D9"),
    right=Side(style="thin", color="D9D9D9"),
    top=Side(style="thin", color="D9D9D9"),
    bottom=Side(style="thin", color="D9D9D9"),
)
TOP_WRAP = Alignment(vertical="top", wrap_text=True)
CURRENCY_FORMAT = '"$"#,##0'
PERCENT_FORMAT = "0.0%"
MAX_COL_WIDTH = 48


PROJECT_COLUMNS = [
    ("Record ID", "record_id"),
    ("Project ID", "project_id"),
    ("Project Name", "project_name"),
    ("Alternative Names", "alternative_project_names"),
    ("Owning Organization", "owning_organization"),
    ("Department / Agency", "department_or_agency"),
    ("Category", "project_category"),
    ("Subcategory", "project_subcategory"),
    ("Infrastructure Domains", "infrastructure_domains"),
    ("Location", "project_location"),
    ("Address / Site", "address_or_site"),
    ("City", "city"),
    ("County", "county"),
    ("State", "state"),
    ("Description", "project_description"),
    ("Problem / Need", "problem_or_need"),
    ("Proposed Scope", "proposed_scope"),
    ("Major Components", "major_components"),
    ("Current Phase", "current_phase"),
    ("Phase Evidence", "phase_evidence"),
    ("Phase Inferred", "phase_is_inferred"),
    ("Pre-RFP Status", "pre_rfp_status"),
    ("Procurement Status", "procurement_status"),
    ("Design Status", "design_status"),
    ("Construction Status", "construction_status"),
    ("Approval Status", "approval_status"),
    ("Funding Status", "funding_status"),
    ("Priority", "priority_level"),
    ("Total Project Budget", "total_project_budget"),
    ("Total Budget Raw", "total_project_budget_raw"),
    ("Current-Year Budget", "current_year_budget"),
    ("Prior-Year Spending", "prior_year_spending"),
    ("Future-Year Funding", "future_year_funding"),
    ("Remaining Budget", "remaining_budget"),
    ("Requested Funding", "requested_funding"),
    ("Approved Funding", "approved_funding"),
    ("Unfunded Amount", "unfunded_amount"),
    ("Bond Funding", "bond_funding"),
    ("Grant Funding", "grant_funding"),
    ("Local Funding", "local_funding"),
    ("State Funding", "state_funding"),
    ("Federal Funding", "federal_funding"),
    ("Other Funding", "other_funding"),
    ("Fund / Account", "fund_or_account"),
    ("Estimated Start", "estimated_start_date"),
    ("Estimated Design", "estimated_design_date"),
    ("Estimated Procurement", "estimated_procurement_date"),
    ("Estimated Bid", "estimated_bid_date"),
    ("Construction Start", "estimated_construction_start"),
    ("Estimated Completion", "estimated_completion_date"),
    ("Timeline Notes", "timeline_notes"),
    ("Approval Body", "approval_body"),
    ("Approval Date", "approval_date"),
    ("Consultants / Engineers", "consultants_or_engineers"),
    ("Contractors / Vendors", "contractors_or_vendors"),
    ("Dependencies", "dependencies"),
    ("Related Projects", "related_projects"),
    ("Risks / Constraints", "risks_or_constraints"),
    ("Source Pages", "source_pages"),
    ("Source Sections", "source_sections"),
    ("Confidence", "confidence_score"),
    ("Quality Flags", "quality_flags"),
    ("Extraction Notes", "extraction_notes"),
]


def export_excel_workbook(
    output: FinalExtractionOutput,
    output_dir: str | Path,
    *,
    source_filename: str,
    run_log: list[RunLogEntry] | list[dict[str, Any]] | None = None,
) -> Path:
    output_dir = ensure_dir(output_dir)
    stem = document_stem(source_filename)
    path = output_dir / f"{stem}_capital_projects.xlsx"

    wb = Workbook()
    # Remove default sheet; recreate in required order.
    default = wb.active
    wb.remove(default)

    _build_summary_sheet(wb, output)
    _build_projects_sheet(wb, output)
    _build_annual_funding_sheet(wb, output)
    _build_funding_sources_sheet(wb, output)
    _build_contacts_sheet(wb, output)
    _build_evidence_sheet(wb, output)
    _build_budget_conflicts_sheet(wb, output)
    _build_validation_issues_sheet(wb, output)
    _build_duplicate_audit_sheet(wb, output)
    _build_run_log_sheet(wb, run_log or [])

    wb.properties.title = f"Capital Projects - {stem}"
    wb.properties.creator = "budget-project-extractor"
    wb.properties.created = datetime.now(timezone.utc).replace(tzinfo=None)
    wb.properties.modified = datetime.now(timezone.utc).replace(tzinfo=None)

    try:
        wb.save(path)
    except PermissionError as exc:
        raise PermissionError(
            f"Unable to write Excel file (is it open in another program?): {path}"
        ) from exc
    return path


def _build_summary_sheet(wb: Workbook, output: FinalExtractionOutput) -> None:
    ws = wb.create_sheet("Summary", 0)
    meta = output.document_metadata
    summary = output.summary
    stats = output.run_statistics

    ws["A1"] = "Capital Project Extraction Summary"
    ws["A1"].font = Font(bold=True, size=16, color="1F4E79")
    ws.merge_cells("A1:D1")

    rows: list[tuple[str, Any]] = [
        ("Document Title", meta.document_title),
        ("Organization", meta.organization_name),
        ("Fiscal Year / Period", meta.fiscal_year_or_period),
        ("Source File", meta.source_file),
        ("Source URL", meta.source_url or ""),
        ("SHA-256", meta.file_sha256),
        ("Total Pages", meta.total_pages),
        ("Processed Pages", meta.processed_pages),
        ("OCR Pages", flatten_list(meta.ocr_pages)),
        ("Failed Pages", flatten_list(meta.failed_pages)),
        ("Processing Started", meta.processing_started_at),
        ("Processing Completed", meta.processing_completed_at),
        ("Model", meta.model),
        ("Application Version", meta.application_version),
        ("Aggregation Method", meta.aggregation_method),
        ("", ""),
        ("Raw Project Records", summary.raw_project_records),
        ("Final Project Count", summary.final_project_count),
        ("Duplicate Records Merged", summary.duplicate_records_merged),
        ("Projects With Known Total Budget", summary.projects_with_known_total_budget),
        ("Projects With Unknown Budget", summary.projects_with_unknown_budget),
        ("Known Total Project Value", summary.known_total_project_value),
        ("Strong Pre-RFP Count", summary.strong_pre_rfp_count),
        ("Possible Pre-RFP Count", summary.possible_pre_rfp_count),
        ("Active Procurement Count", summary.active_procurement_count),
        ("Unknown Phase Count", summary.unknown_phase_count),
        ("Budget Conflict Count", summary.budget_conflict_count),
        ("Low Confidence Count", summary.low_confidence_count),
        ("", ""),
        ("Chunks Total", stats.chunks_total),
        ("Chunks Completed", stats.chunks_completed),
        ("Chunks Failed", stats.chunks_failed),
        ("API Requests", stats.api_requests),
        ("API Retries", stats.api_retries),
        ("Prompt Tokens", stats.prompt_tokens),
        ("Completion Tokens", stats.completion_tokens),
        ("Total Tokens", stats.total_tokens),
        ("Pages Native Text", stats.pages_native_text),
        ("Pages OCR", stats.pages_ocr),
        ("Pages Failed", stats.pages_failed),
    ]

    ws.append(["Field", "Value"])
    _style_header(ws, 2, 2)
    for label, value in rows:
        ws.append([label, value])

    # Currency for known total
    for row_idx in range(3, ws.max_row + 1):
        if ws.cell(row_idx, 1).value == "Known Total Project Value":
            ws.cell(row_idx, 2).number_format = CURRENCY_FORMAT

    # Distributions
    start = ws.max_row + 2
    ws.cell(start, 1, "Phase Distribution (budget sum)")
    ws.cell(start, 1).font = Font(bold=True)
    ws.append(["Phase", "Budget Sum"])
    phase_header_row = ws.max_row
    _style_header(ws, phase_header_row, 2)
    for key, value in summary.sums_by_phase.items():
        ws.append([key, value])
        ws.cell(ws.max_row, 2).number_format = CURRENCY_FORMAT

    ws.append([])
    ws.append(["Category", "Budget Sum"])
    _style_header(ws, ws.max_row, 2)
    for key, value in summary.sums_by_category.items():
        ws.append([key, value])
        ws.cell(ws.max_row, 2).number_format = CURRENCY_FORMAT

    ws.append([])
    ws.append(["Department", "Budget Sum"])
    _style_header(ws, ws.max_row, 2)
    for key, value in summary.sums_by_department.items():
        ws.append([key, value])
        ws.cell(ws.max_row, 2).number_format = CURRENCY_FORMAT

    legend_row = ws.max_row + 2
    ws.cell(legend_row, 1, "Legend")
    ws.cell(legend_row, 1).font = Font(bold=True)
    ws.cell(legend_row + 1, 1, "Green")
    ws.cell(legend_row + 1, 2, "Strong pre-RFP project")
    ws.cell(legend_row + 1, 1).fill = GREEN_FILL
    ws.cell(legend_row + 2, 1, "Yellow")
    ws.cell(legend_row + 2, 2, "Uncertain / incomplete / low-confidence")
    ws.cell(legend_row + 2, 1).fill = YELLOW_FILL
    ws.cell(legend_row + 3, 1, "Red")
    ws.cell(legend_row + 3, 2, "Failed validation / error severity")
    ws.cell(legend_row + 3, 1).fill = RED_FILL

    if meta.source_url:
        # Find source URL row and hyperlink it.
        for row_idx in range(3, 40):
            if ws.cell(row_idx, 1).value == "Source URL" and ws.cell(row_idx, 2).value:
                ws.cell(row_idx, 2).hyperlink = str(ws.cell(row_idx, 2).value)
                ws.cell(row_idx, 2).font = Font(color="0563C1", underline="single")
                break

    _finalize_sheet(ws, freeze_row=2)


def _build_projects_sheet(wb: Workbook, output: FinalExtractionOutput) -> None:
    ws = wb.create_sheet("Projects")
    headers = [col[0] for col in PROJECT_COLUMNS]
    ws.append(headers)
    _style_header(ws, 1, len(headers))

    currency_cols = {
        "Total Project Budget",
        "Current-Year Budget",
        "Prior-Year Spending",
        "Future-Year Funding",
        "Remaining Budget",
        "Requested Funding",
        "Approved Funding",
        "Unfunded Amount",
        "Bond Funding",
        "Grant Funding",
        "Local Funding",
        "State Funding",
        "Federal Funding",
        "Other Funding",
    }
    header_index = {name: idx + 1 for idx, name in enumerate(headers)}

    for project in output.projects:
        data = project.model_dump(mode="json")
        row = []
        for _, key in PROJECT_COLUMNS:
            value = data.get(key)
            if isinstance(value, list):
                if key == "source_pages":
                    value = flatten_list(value, sep=", ")
                else:
                    value = flatten_list(value, sep="; ")
            if key in {"current_phase", "pre_rfp_status"} and value is not None:
                value = str(value)
            row.append(value)
        ws.append(row)
        row_idx = ws.max_row
        for name in currency_cols:
            ws.cell(row_idx, header_index[name]).number_format = CURRENCY_FORMAT
        ws.cell(row_idx, header_index["Confidence"]).number_format = PERCENT_FORMAT
        # Confidence is 0-1 already.
        pre_rfp = str(data.get("pre_rfp_status") or "")
        flags = flatten_list(data.get("quality_flags") or [])
        if pre_rfp == "strong_pre_rfp":
            ws.cell(row_idx, header_index["Pre-RFP Status"]).fill = GREEN_FILL
        if "incomplete_chunk_boundary" in flags or "low_confidence" in flags or "possible_duplicate" in flags:
            ws.cell(row_idx, header_index["Quality Flags"]).fill = YELLOW_FILL

    # Conditional formatting for pre-RFP / phase helpers
    pre_col = get_column_letter(header_index["Pre-RFP Status"])
    ws.conditional_formatting.add(
        f"{pre_col}2:{pre_col}{max(ws.max_row, 2)}",
        CellIsRule(operator="equal", formula=['"strong_pre_rfp"'], fill=GREEN_FILL),
    )
    ws.conditional_formatting.add(
        f"{pre_col}2:{pre_col}{max(ws.max_row, 2)}",
        CellIsRule(operator="equal", formula=['"insufficient_information"'], fill=YELLOW_FILL),
    )

    _add_table(ws, "ProjectsTable")
    _finalize_sheet(ws)


def _build_annual_funding_sheet(wb: Workbook, output: FinalExtractionOutput) -> None:
    ws = wb.create_sheet("Annual Funding")
    headers = [
        "Record ID",
        "Project ID",
        "Project Name",
        "Fiscal Year",
        "Amount",
        "Raw Amount",
        "Funding Status",
        "Source Pages",
    ]
    ws.append(headers)
    _style_header(ws, 1, len(headers))
    for project in output.projects:
        for row in project.annual_funding_schedule:
            ws.append(
                [
                    project.record_id,
                    project.project_id,
                    project.project_name,
                    row.fiscal_year,
                    row.amount,
                    row.amount_raw,
                    row.funding_status,
                    flatten_list(row.source_pages, sep=", "),
                ]
            )
            ws.cell(ws.max_row, 5).number_format = CURRENCY_FORMAT
    _add_table(ws, "AnnualFundingTable")
    _finalize_sheet(ws)


def _build_funding_sources_sheet(wb: Workbook, output: FinalExtractionOutput) -> None:
    ws = wb.create_sheet("Funding Sources")
    headers = [
        "Record ID",
        "Project ID",
        "Project Name",
        "Funding Source",
        "Source Type",
        "Amount",
        "Raw Amount",
        "Fiscal Year",
        "Source Pages",
    ]
    ws.append(headers)
    _style_header(ws, 1, len(headers))
    for project in output.projects:
        for row in project.funding_sources:
            ws.append(
                [
                    project.record_id,
                    project.project_id,
                    project.project_name,
                    row.source_name,
                    row.source_type,
                    row.amount,
                    row.amount_raw,
                    row.fiscal_year,
                    flatten_list(row.source_pages, sep=", "),
                ]
            )
            ws.cell(ws.max_row, 6).number_format = CURRENCY_FORMAT
    _add_table(ws, "FundingSourcesTable")
    _finalize_sheet(ws)


def _build_contacts_sheet(wb: Workbook, output: FinalExtractionOutput) -> None:
    ws = wb.create_sheet("Contacts")
    headers = [
        "Record ID",
        "Project ID",
        "Project Name",
        "Contact Name",
        "Title",
        "Department",
        "Email",
        "Phone",
        "Source Pages",
    ]
    ws.append(headers)
    _style_header(ws, 1, len(headers))
    for project in output.projects:
        for contact in project.responsible_contacts:
            ws.append(
                [
                    project.record_id,
                    project.project_id,
                    project.project_name,
                    contact.name,
                    contact.title,
                    contact.department,
                    contact.email,
                    contact.phone,
                    flatten_list(contact.source_pages, sep=", "),
                ]
            )
    _add_table(ws, "ContactsTable")
    _finalize_sheet(ws)


def _build_evidence_sheet(wb: Workbook, output: FinalExtractionOutput) -> None:
    ws = wb.create_sheet("Evidence")
    headers = [
        "Record ID",
        "Project ID",
        "Project Name",
        "PDF Page",
        "Evidence Type",
        "Evidence Quote",
    ]
    ws.append(headers)
    _style_header(ws, 1, len(headers))
    for project in output.projects:
        for evidence in project.evidence:
            ws.append(
                [
                    project.record_id,
                    project.project_id,
                    project.project_name,
                    evidence.page,
                    evidence.evidence_type,
                    evidence.quote,
                ]
            )
    _add_table(ws, "EvidenceTable")
    _finalize_sheet(ws, width_overrides={6: 60})


def _build_budget_conflicts_sheet(wb: Workbook, output: FinalExtractionOutput) -> None:
    ws = wb.create_sheet("Budget Conflicts")
    headers = [
        "Record ID",
        "Project ID",
        "Project Name",
        "Field",
        "Conflicting Values",
        "Source Pages",
        "Explanation",
    ]
    ws.append(headers)
    _style_header(ws, 1, len(headers))
    for project in output.projects:
        for conflict in project.budget_conflicts:
            ws.append(
                [
                    project.record_id,
                    project.project_id,
                    project.project_name,
                    conflict.field_name,
                    flatten_list(conflict.values),
                    flatten_list(conflict.source_pages, sep=", "),
                    conflict.explanation,
                ]
            )
            for col in range(1, 8):
                ws.cell(ws.max_row, col).fill = YELLOW_FILL
    _add_table(ws, "BudgetConflictsTable")
    _finalize_sheet(ws)


def _build_validation_issues_sheet(wb: Workbook, output: FinalExtractionOutput) -> None:
    ws = wb.create_sheet("Validation Issues")
    headers = [
        "Severity",
        "Issue Type",
        "Record ID",
        "Project Name",
        "Chunk ID",
        "Pages",
        "Description",
        "Recommended Review",
    ]
    ws.append(headers)
    _style_header(ws, 1, len(headers))
    for issue in output.validation_issues:
        ws.append(
            [
                issue.severity,
                issue.issue_type,
                issue.record_id,
                issue.project_name,
                issue.chunk_id,
                flatten_list(issue.pages, sep=", "),
                issue.description,
                issue.recommended_review,
            ]
        )
        if issue.severity == "error":
            for col in range(1, 9):
                ws.cell(ws.max_row, col).fill = RED_FILL
        elif issue.severity == "warning":
            for col in range(1, 9):
                ws.cell(ws.max_row, col).fill = YELLOW_FILL
    _add_table(ws, "ValidationIssuesTable")
    _finalize_sheet(ws)


def _build_duplicate_audit_sheet(wb: Workbook, output: FinalExtractionOutput) -> None:
    ws = wb.create_sheet("Duplicate Audit")
    headers = [
        "Decision",
        "Final Record ID",
        "Candidate Record IDs",
        "Candidate Names",
        "Similarity Scores",
        "Merge Reason",
        "Source Pages",
        "Resolution Method",
        "Resolution Confidence",
    ]
    ws.append(headers)
    _style_header(ws, 1, len(headers))
    for entry in output.duplicate_audit:
        ws.append(
            [
                entry.decision,
                entry.final_record_id,
                flatten_list(entry.candidate_record_ids),
                flatten_list(entry.candidate_names),
                flatten_list(entry.similarity_scores),
                entry.merge_reason,
                flatten_list(entry.source_pages, sep=", "),
                entry.resolution_method,
                entry.resolution_confidence,
            ]
        )
        if entry.decision == "uncertain":
            for col in range(1, 10):
                ws.cell(ws.max_row, col).fill = YELLOW_FILL
    _add_table(ws, "DuplicateAuditTable")
    _finalize_sheet(ws)


def _build_run_log_sheet(
    wb: Workbook,
    run_log: list[RunLogEntry] | list[dict[str, Any]],
) -> None:
    ws = wb.create_sheet("Run Log")
    headers = [
        "Timestamp",
        "Level",
        "Stage",
        "Chunk ID",
        "Page Range",
        "Request ID",
        "Model",
        "Attempt",
        "Status",
        "Prompt Tokens",
        "Completion Tokens",
        "Message",
    ]
    ws.append(headers)
    _style_header(ws, 1, len(headers))
    for entry in run_log:
        data = entry.__dict__ if hasattr(entry, "__dict__") and not isinstance(entry, dict) else entry
        ws.append(
            [
                data.get("timestamp"),
                data.get("level"),
                data.get("stage"),
                data.get("chunk_id"),
                data.get("page_range"),
                data.get("request_id"),
                data.get("model"),
                data.get("attempt"),
                data.get("status"),
                data.get("prompt_tokens"),
                data.get("completion_tokens"),
                data.get("message"),
            ]
        )
        if str(data.get("level", "")).upper() == "ERROR":
            for col in range(1, 13):
                ws.cell(ws.max_row, col).fill = RED_FILL
    _add_table(ws, "RunLogTable")
    _finalize_sheet(ws, width_overrides={12: 60})


def _style_header(ws, row: int, col_count: int) -> None:
    for col in range(1, col_count + 1):
        cell = ws.cell(row, col)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(vertical="center", wrap_text=True)
        cell.border = THIN


def _add_table(ws, name: str) -> None:
    if ws.max_row < 1 or ws.max_column < 1:
        return
    ref = f"A1:{get_column_letter(ws.max_column)}{max(ws.max_row, 1)}"
    # Table names must be unique and not contain spaces issues; sanitize.
    safe = "".join(ch for ch in name if ch.isalnum() or ch == "_")
    table = Table(displayName=safe, ref=ref)
    table.tableStyleInfo = TableStyleInfo(
        name="TableStyleMedium2",
        showFirstColumn=False,
        showLastColumn=False,
        showRowStripes=True,
        showColumnStripes=False,
    )
    # Avoid crashing when sheet has only headers.
    try:
        ws.add_table(table)
    except Exception:
        # Fallback: apply alternating fills manually.
        for row in range(2, ws.max_row + 1):
            if row % 2 == 0:
                for col in range(1, ws.max_column + 1):
                    if ws.cell(row, col).fill.fgColor is None or ws.cell(row, col).fill.fill_type is None:
                        ws.cell(row, col).fill = ALT_ROW_FILL


def _finalize_sheet(
    ws,
    *,
    freeze_row: int = 1,
    width_overrides: dict[int, int] | None = None,
) -> None:
    ws.freeze_panes = f"A{freeze_row + 1}"
    ws.auto_filter.ref = f"A{freeze_row}:{get_column_letter(ws.max_column)}{max(ws.max_row, freeze_row)}"
    width_overrides = width_overrides or {}

    for col in range(1, ws.max_column + 1):
        letter = get_column_letter(col)
        max_len = 0
        for row in range(1, min(ws.max_row, 200) + 1):
            value = ws.cell(row, col).value
            if value is None:
                continue
            max_len = max(max_len, len(str(value)))
        width = min(MAX_COL_WIDTH, max(10, max_len + 2))
        if col in width_overrides:
            width = width_overrides[col]
        ws.column_dimensions[letter].width = width

    for row in ws.iter_rows(min_row=1, max_row=ws.max_row, max_col=ws.max_column):
        for cell in row:
            cell.alignment = TOP_WRAP
            if cell.border.left.style is None:
                cell.border = THIN
