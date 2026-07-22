from __future__ import annotations

from pathlib import Path

from openpyxl import load_workbook

from budget_extractor.excel_exporter import export_excel_workbook
from budget_extractor.schemas import (
    DocumentMetadata,
    ExtractionSummary,
    FinalExtractionOutput,
    ProjectRecord,
    RunStatistics,
    ValidationIssue,
)


REQUIRED_SHEETS = [
    "Summary",
    "Projects",
    "Annual Funding",
    "Funding Sources",
    "Contacts",
    "Evidence",
    "Budget Conflicts",
    "Validation Issues",
    "Duplicate Audit",
    "Run Log",
]


def test_excel_workbook_creation(tmp_path: Path):
    project = ProjectRecord.model_validate(
        {
            "record_id": "proj_test",
            "project_name": "City Hall Roof Replacement",
            "current_phase": "rfq_rfp_preparation",
            "pre_rfp_status": "strong_pre_rfp",
            "source_pages": [5],
            "evidence": [{"page": 5, "quote": "Roof replacement planned for FY26"}],
            "confidence_score": 0.91,
            "first_seen_chunk": "chunk_0001",
            "total_project_budget": 850000,
            "department_or_agency": "Facilities",
            "annual_funding_schedule": [
                {
                    "fiscal_year": "FY2026",
                    "amount": 850000,
                    "amount_raw": "$850,000",
                    "source_pages": [5],
                }
            ],
        }
    )
    output = FinalExtractionOutput(
        document_metadata=DocumentMetadata(
            document_title="Test Budget",
            organization_name="Test City",
            source_file="test_budget.pdf",
            source_url="https://example.gov/test_budget.pdf",
            file_sha256="abc",
            total_pages=10,
            processed_pages=10,
            model="glm-5.2",
            application_version="1.0.0",
            aggregation_method="test",
        ),
        summary=ExtractionSummary(
            raw_project_records=1,
            final_project_count=1,
            projects_with_known_total_budget=1,
            known_total_project_value=850000,
            strong_pre_rfp_count=1,
        ),
        projects=[project],
        validation_issues=[
            ValidationIssue(
                severity="error",
                issue_type="failed_chunk",
                chunk_id="chunk_0009",
                description="Chunk failed validation",
                recommended_review="Re-run chunk",
            )
        ],
        run_statistics=RunStatistics(chunks_total=1, chunks_completed=1),
    )

    path = export_excel_workbook(output, tmp_path, source_filename="test_budget.pdf", run_log=[])
    assert path.exists()

    wb = load_workbook(path)
    assert wb.sheetnames == REQUIRED_SHEETS
    assert wb["Projects"].cell(2, 3).value == "City Hall Roof Replacement"
    assert wb["Annual Funding"].cell(2, 4).value == "FY2026"
    assert wb["Summary"].cell(1, 1).value
