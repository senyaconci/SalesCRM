from pathlib import Path

from org_intel.config import RunConfig, Settings
from org_intel.database import Database
from org_intel.llm.glm_provider import MockLLMProvider
from org_intel.pipeline.orchestrator import PipelineOrchestrator
from org_intel.retrieval.cache import ContentCache
from org_intel.schemas.enums import OrganizationType


def _run(org_url: str, http, tmp_path: Path, mode: str = "full") -> PipelineOrchestrator:
    settings = Settings(
        llm_provider="mock",
        cache_dir=tmp_path / "cache",
        max_cost_usd=10.0,
        respect_robots=False,
        official_sources_only=True,
        max_documents=50,
    )
    run = RunConfig(
        org_url=org_url,
        output_dir=tmp_path / "out",
        mode=mode,
        include_completed=True,
    )
    orch = PipelineOrchestrator(
        settings,
        run,
        http=http,
        llm_provider=MockLLMProvider(),
    )
    orch.run_pipeline()
    return orch


def test_municipality_full_pipeline(municipality_http, tmp_path: Path):
    orch = _run("https://www.examplecity.gov", municipality_http, tmp_path)
    assert orch.state.identity is not None
    assert orch.state.identity.organization_type == OrganizationType.CITY
    assert orch.state.source_registry and orch.state.source_registry.sources
    assert orch.state.document_inventory and orch.state.document_inventory.documents
    assert orch.state.anchor_selection.get("document")
    # Prefer CIP/html/csv over giant budget PDF
    reason = orch.state.anchor_selection.get("reason", "").lower()
    assert "pdf" not in reason or "registry" in reason or "cip" in reason or "csv" in reason or "live" in reason
    assert orch.state.project_index
    ids = {p.project_id for p in orch.state.project_index}
    assert "W0280" in ids
    assert (tmp_path / "out" / "organization_profile.json").exists()
    assert (tmp_path / "out" / "organization_intelligence.md").exists()
    assert (tmp_path / "out" / "capital_projects.xlsx").exists()
    assert (tmp_path / "out" / "opportunities.json").exists()
    assert (tmp_path / "out" / "cost_ledger.json").exists()
    assert (tmp_path / "out" / "run_manifest.json").exists()


def test_water_university_airport_discover(water_http, university_http, airport_http, tmp_path: Path):
    for url, http, expected in [
        ("https://www.examplewater.gov", water_http, OrganizationType.WATER_DISTRICT),
        ("https://www.exampleu.edu", university_http, OrganizationType.UNIVERSITY),
        ("https://www.exampleairport.gov", airport_http, OrganizationType.AIRPORT),
    ]:
        out = tmp_path / expected.value
        settings = Settings(llm_provider="mock", cache_dir=out / "cache", respect_robots=False)
        run = RunConfig(org_url=url, output_dir=out, mode="discover")
        orch = PipelineOrchestrator(settings, run, http=http, llm_provider=MockLLMProvider())
        orch.run_pipeline()
        assert orch.state.identity is not None
        assert orch.state.identity.organization_type == expected
        assert orch.state.org_graph and orch.state.org_graph.nodes
        assert orch.state.source_registry


def test_resume_behavior(municipality_http, tmp_path: Path):
    orch = _run("https://www.examplecity.gov", municipality_http, tmp_path / "r1", mode="discover")
    run_id = orch.run_id
    # Second run resume full should reuse identity checkpoint
    settings = Settings(llm_provider="mock", cache_dir=tmp_path / "r1" / "cache", respect_robots=False)
    run = RunConfig(
        org_url="https://www.examplecity.gov",
        output_dir=tmp_path / "r1" / "out",
        mode="full",
        resume=True,
    )
    orch2 = PipelineOrchestrator(
        settings, run, http=municipality_http, llm_provider=MockLLMProvider(), run_id=run_id
    )
    orch2.run_pipeline()
    assert orch2.state.identity is not None
    assert "canonical_identity" in orch2.db.completed_phases(run_id)


def test_cache_behavior(municipality_http, tmp_path: Path):
    first = municipality_http.fetch("https://www.examplecity.gov")
    assert first.from_cache is False
    second = municipality_http.fetch("https://www.examplecity.gov")
    assert second.from_cache is True
    assert first.sha256 == second.sha256


def test_cost_limit_enforcement(municipality_http, tmp_path: Path, settings, tmp_db):
    from org_intel.llm.router import CostLimitExceeded, CostTracker, LLMRouter
    from org_intel.llm.base import ModelRole
    from org_intel.schemas.evidence import CostLedgerEntry

    tracker = CostTracker(tmp_db, "RUN-COST", max_cost_usd=0.0001)
    router = LLMRouter(settings, tmp_db, tracker, provider=MockLLMProvider())
    # Force a non-cache expensive-looking call by disabling cache and tiny limit
    router.cache_enabled = False
    # Manually record near-limit cost then fail
    tracker.record(
        CostLedgerEntry(
            task_type="setup",
            estimated_cost=0.0001,
            actual_cost=0.0001,
        )
    )
    try:
        router.complete(
            ModelRole.REASONING_MODEL,
            "sys",
            "user " * 5000,
            task_type="test",
        )
        assert False, "expected CostLimitExceeded"
    except CostLimitExceeded:
        pass


def test_output_schemas(municipality_http, tmp_path: Path):
    orch = _run("https://www.examplecity.gov", municipality_http, tmp_path)
    import json

    profile = json.loads((tmp_path / "out" / "organization_profile.json").read_text())
    assert "canonical_name" in profile
    assert "organization_id" in profile
    projects = json.loads((tmp_path / "out" / "projects_full.json").read_text())
    assert isinstance(projects, list)
    for p in projects:
        assert "record_id" in p
        assert "pre_rfq_classification" in p
        assert "evidence" in p


def test_excel_generation(municipality_http, tmp_path: Path):
    from openpyxl import load_workbook

    _run("https://www.examplecity.gov", municipality_http, tmp_path)
    wb = load_workbook(tmp_path / "out" / "capital_projects.xlsx")
    expected_sheets = {
        "Executive Summary",
        "Organization",
        "Project Index",
        "Full Projects",
        "Opportunities",
        "Cost Ledger",
        "Research Gaps",
    }
    assert expected_sheets.issubset(set(wb.sheetnames))


def test_incremental_refresh(municipality_http, tmp_path: Path):
    orch = _run("https://www.examplecity.gov", municipality_http, tmp_path / "ref")
    settings = Settings(llm_provider="mock", cache_dir=tmp_path / "ref" / "cache", respect_robots=False)
    run = RunConfig(
        org_url="https://www.examplecity.gov",
        output_dir=tmp_path / "ref" / "out",
        mode="refresh",
        resume=True,
    )
    orch2 = PipelineOrchestrator(
        settings,
        run,
        http=municipality_http,
        llm_provider=MockLLMProvider(),
        run_id=orch.run_id,
    )
    orch2.run_pipeline()
    assert (tmp_path / "ref" / "out" / "change_history.json").exists()


def test_conflicting_evidence_preserved():
    from org_intel.projects.consolidation import consolidate_project
    from org_intel.schemas.project import ProjectRecord
    from org_intel.schemas.enums import ProjectPhase

    project = ProjectRecord(
        organization_id="ORG",
        project_name="Test",
        total_project_cost=1000,
        current_year_appropriation=1000,
        normalized_phase=ProjectPhase.FUNDING_APPROVED,
        consultants=[],
    )
    # Add consultant to create phase conflict
    from org_intel.schemas.project import CompanyRelationship

    project.consultants.append(CompanyRelationship(company_name="Acme Engineering"))
    out = consolidate_project(project)
    assert out.budget_conflicts or out.quality_flags
    assert out.source_conflicts
